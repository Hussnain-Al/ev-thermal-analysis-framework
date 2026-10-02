"""Independent Python cross-check of the literature gap-fill calculations.

The MATLAB module ``modules/literature_gap_fill`` is the project
implementation. This script re-implements the same equations with the same
inputs so the numbers can be checked without MATLAB, and renders the preview
figures stored in ``docs/images/gap_fill``.

Run from the repository root:

    python tools/gap_fill_reference.py

Requires numpy, scipy, matplotlib and openpyxl.
"""

from __future__ import annotations

import csv
import math
from pathlib import Path

import numpy as np
import openpyxl
from scipy.interpolate import RegularGridInterpolator

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
FIG_DIR = ROOT / "docs" / "images" / "gap_fill"


def load_assumptions() -> dict[str, float]:
    path = ROOT / "data" / "literature" / "literature_assumption_register.csv"
    with path.open(newline="", encoding="utf-8") as fh:
        return {row["ID"]: float(row["Central"]) for row in csv.DictReader(fh)}


A = load_assumptions()

# ---------------------------------------------------------------------------
# Project inputs that already exist in the MATLAB configuration.
# ---------------------------------------------------------------------------
VEHICLE = dict(mass=1950.0, g=9.81, r=0.724 / 2, gear=9.11,
               A=566.4645, B=0.4185918)
CAPACITY_AH = 134.0
SERIES = 108
NOMINAL_V = 3.2
ACR = 0.40e-3
R_BASE_RECON = 3.10
LIMIT_CHARGE = 55.0
LIMIT_ABS = 60.0
RAD_IN = 65.0
AIR_IN = 45.0
DESIGN_FLOW_LMIN = 20.0
COOLANT_60 = dict(rho=1040.0, mu=1.50e-3, cp=3600.0)
DESIGN_DUTY_KW = {"Sustained 10% grade": 2.769, "Low-speed hot-weather grade": 1.549}


# ---------------------------------------------------------------------------
# Drive unit: operating points on the supplied efficiency surface.
# ---------------------------------------------------------------------------
def load_map():
    rows = list(csv.DictReader((ROOT / "data/motor_heat/drive_unit_efficiency_map.csv").open()))
    rpm = sorted({float(r["Speed_rpm"]) for r in rows})
    tq = sorted({float(r["Torque_Nm"]) for r in rows})
    eta = np.zeros((len(rpm), len(tq)))
    for r in rows:
        eta[rpm.index(float(r["Speed_rpm"])), tq.index(float(r["Torque_Nm"]))] = float(r["IntegratedEfficiency_pct"]) / 100
    wb = openpyxl.load_workbook(ROOT / "data/motor_heat/drive_unit_limits.xlsx", data_only=True)
    t = np.array([r[:2] for r in wb["Torque RPM Curve"].iter_rows(min_row=2, values_only=True) if r[0] is not None], float)
    p = np.array([r[:2] for r in wb["Power RPM Curve"].iter_rows(min_row=2, values_only=True) if r[0] is not None], float)
    t = t[np.argsort(t[:, 0])]
    p = p[np.argsort(p[:, 0])]
    return np.array(rpm), np.array(tq), eta, t, p


def read_cycle(name):
    vals = []
    for line in (ROOT / "data/common/cycles" / name).read_text().splitlines():
        parts = line.split()
        try:
            vals.append((float(parts[0]), float(parts[1])))
        except (ValueError, IndexError):
            continue
    arr = np.array(vals)
    return arr[:, 0], arr[:, 1] * 0.44704


def operating_trace(t, v, grade_pct, mp):
    rpm_ax, tq_ax, eta, _, _ = mp
    interp = RegularGridInterpolator((rpm_ax, tq_ax), eta)
    a = np.gradient(v, t)
    omega = v / VEHICLE["r"] * VEHICLE["gear"]
    rpm = omega * 60 / (2 * np.pi)
    force = VEHICLE["A"] + VEHICLE["B"] * v**2 + VEHICLE["mass"] * a + \
        VEHICLE["mass"] * VEHICLE["g"] * math.sin(math.atan(grade_pct / 100))
    p_wheel = force * v / 1000
    torque = np.zeros_like(t)
    moving = omega > 1e-6
    torque[moving] = p_wheel[moving] * 1000 / omega[moving]
    q = np.column_stack([np.clip(np.abs(rpm), rpm_ax[0], rpm_ax[-1]),
                         np.clip(np.abs(torque), tq_ax[0], tq_ax[-1])])
    e = interp(q)
    e[~moving] = 1
    dc = np.where(p_wheel >= 0, p_wheel / e, p_wheel * e)
    heat = np.where(p_wheel >= 0, dc - p_wheel, np.abs(p_wheel) * (1 - e))
    return dict(t=t, rpm=rpm, torque=torque, p_wheel=p_wheel, dc=dc, heat=heat, eta=e)


def drive_traces(mp):
    out = {}
    for label, f in [("Urban stop-start", "urban_cycle.txt"), ("Highway", "highway_cycle.txt")]:
        t, v = read_cycle(f)
        out[label] = operating_trace(t, v, 0.0, mp)
    for label, kmh, grade, dur in [("Sustained 10% grade", 40, 10, 1200),
                                   ("Low-speed hot-weather grade", 15, 5, 1800)]:
        t = np.arange(0, dur + 1, dtype=float)
        v = np.full_like(t, kmh / 1.609344 * 0.44704)
        out[label] = operating_trace(t, v, grade, mp)
    return out


# ---------------------------------------------------------------------------
# Drive unit: winding-to-coolant resistance implied by the supplier reference.
# ---------------------------------------------------------------------------
def implied_winding_resistance(mp):
    rpm_ax, tq_ax, eta, _, _ = mp
    interp = RegularGridInterpolator((rpm_ax, tq_ax), eta)
    rated_kw = A["M01"]
    controller_loss_kw = 1.580
    speeds = np.arange(3000, 9001, 250.0)
    torque = rated_kw * 1000 / (speeds * 2 * np.pi / 60)
    e = interp(np.column_stack([speeds, torque]))
    integrated_loss = rated_kw / e - rated_kw
    motor_loss = integrated_loss - controller_loss_kw
    rise = 143.0 - 60.0
    r_hot = rise / (motor_loss * 1000)
    return speeds, torque, e, integrated_loss, motor_loss, r_hot


# ---------------------------------------------------------------------------
# Radiator: e-NTU with Chang and Wang (1997) louvered-fin j-factor.
# ---------------------------------------------------------------------------
def air_props(T):
    rho = 101325 / (287.05 * (T + 273.15))
    mu = 1.716e-5 * ((T + 273.15) / 273.15) ** 1.5 * (273.15 + 110.4) / (T + 273.15 + 110.4)
    return rho, mu, 1007.0, 0.0263 + 7.4e-5 * (T - 27), 0.705


def radiator_performance(face_v, flow_lmin, coolant=COOLANT_60, t_cool_in=RAD_IN, t_air_in=AIR_IN):
    # Candidate geometry from data/motor_cooling/propulsion_radiator_geometry.csv
    n_tubes, tube_len, td, tube_h = 31, 0.275, 0.026, 0.002
    wall, fin_t, fin_h, fin_p, fin_len, n_channels = 0.2e-3, 0.1e-3, 0.008, 2.8e-3, 0.260, 32
    a_front = 0.0837
    lp, theta, ll = A["R01"] * 1e-3, A["R02"], A["R03"] * fin_h
    k_fin = A["R04"]
    tp = fin_h + tube_h

    rho, mu, cp_a, k_a, pr = air_props(t_air_in)
    # Air-side areas
    fins_per_channel = fin_len / fin_p
    a_fin = n_channels * fins_per_channel * 2 * fin_h * td
    a_tube = n_tubes * 2 * td * tube_len * (1 - fin_t / fin_p)
    a_air = a_fin + a_tube
    a_free = n_channels * fin_len * fin_h * (1 - fin_t / fin_p)
    sigma = a_free / a_front

    v_core = face_v / sigma
    re_lp = rho * v_core * lp / mu
    j = (re_lp ** -0.49 * (theta / 90) ** 0.27 * (fin_p / lp) ** -0.14 * (fin_h / lp) ** -0.29
         * (td / lp) ** -0.23 * (ll / lp) ** 0.68 * (tp / lp) ** -0.28 * (fin_t / lp) ** -0.05)
    h_air = j * rho * v_core * cp_a / pr ** (2 / 3)
    m = np.sqrt(2 * h_air / (k_fin * fin_t))
    lf = fin_h / 2
    eta_fin = np.tanh(m * lf) / (m * lf)
    eta_o = 1 - a_fin / a_air * (1 - eta_fin)

    # Coolant side: laminar flat-tube flow, fully developed Nu for a 16:1 duct.
    wi, hi = td - 2 * wall, tube_h - 2 * wall
    a_x = wi * hi
    dh = 4 * a_x / (2 * (wi + hi))
    q_tube = flow_lmin / 60000 / n_tubes
    v_c = q_tube / a_x
    re_c = coolant["rho"] * v_c * dh / coolant["mu"]
    k_c = A["R05"]
    nu = A["R06"]
    h_c = nu * k_c / dh
    a_c = n_tubes * 2 * (wi + hi) * tube_len

    ua = 1 / (1 / (eta_o * h_air * a_air) + 1 / (h_c * a_c))
    c_air = rho * face_v * a_front * cp_a
    c_cool = flow_lmin / 60000 * coolant["rho"] * coolant["cp"]
    c_min, c_max = np.minimum(c_air, c_cool), np.maximum(c_air, c_cool)
    cr = c_min / c_max
    ntu = ua / c_min
    eps = 1 - np.exp(ntu ** 0.22 / cr * (np.exp(-cr * ntu ** 0.78) - 1))
    q = eps * c_min * (t_cool_in - t_air_in)

    # Coolant pressure drop through tubes (laminar, f*Re for 16:1 rectangle)
    f_re = A["R07"]
    dp_tubes = f_re / re_c * tube_len / dh * coolant["rho"] * v_c**2 / 2
    dp_headers = A["R08"] * coolant["rho"] * v_c**2 / 2
    return dict(q=q, ua=ua, h_air=h_air, eta_fin=eta_fin, re_lp=re_lp, re_c=re_c,
                h_c=h_c, a_air=a_air, a_c=a_c, sigma=sigma, dp=dp_tubes + dp_headers)


# ---------------------------------------------------------------------------
# Battery: DCIR, entropic heat, cell-to-coolant path, discharge transient.
# ---------------------------------------------------------------------------
SOC_TABLE = np.array([0, 5, 20, 37.8, 45, 55, 65.5, 75, 88.5, 95, 100])
DUDT_TABLE = np.array([-0.30, -0.37, -0.15, 0.0, 0.10, 0.05, 0.0, -0.05, 0.0, 0.03, 0.03]) * 1e-3


def dcir(temp_c):
    r25 = ACR / A["B01"]
    ea = A["B02"] * 1000
    return r25 * np.exp(ea / 8.314 * (1 / (temp_c + 273.15) - 1 / 298.15))


def path_budget():
    w, d, h = A["B10"] * 1e-3, A["B11"] * 1e-3, A["B12"] * 1e-3
    base = w * d
    r_internal = h / (3 * A["B13"] * base)
    r_film = A["B14"] * 1e-3 / (A["B15"] * base)
    r_tim = A["B16"] * 1e-3 / (A["B17"] * base)
    r_plate = 1 / (A["B18"] * base * A["B19"])
    return {"Cell interior (mean, axial)": r_internal, "Insulation film + can": r_film,
            "Thermal pad": r_tim, "Cold-plate convection": r_plate}


def battery_discharge(c_rate, t_coolant, r_path, dt=1.0):
    m_cp = A["B20"] * A["B21"]
    i = c_rate * CAPACITY_AH
    dur = 3600 / c_rate
    n = int(dur / dt) + 1
    t = np.arange(n) * dt
    T = np.zeros(n)
    T[0] = t_coolant
    q_j = np.zeros(n)
    q_r = np.zeros(n)
    soc = 100 - 100 * t / dur
    for k in range(n):
        q_j[k] = i**2 * dcir(T[k])
        q_r[k] = -i * (T[k] + 273.15) * np.interp(soc[k], SOC_TABLE, DUDT_TABLE)
        if k < n - 1:
            T[k + 1] = T[k] + dt * (q_j[k] + q_r[k] - (T[k] - t_coolant) / r_path) / m_cp
    return t, T, q_j, q_r, soc


# ---------------------------------------------------------------------------
# Cabin: workbook audit and heat-balance rebuild.
# ---------------------------------------------------------------------------
def psat(T):
    return 610.94 * math.exp(17.625 * T / (T + 243.04))


def humidity_ratio(T, rh):
    pv = rh / 100 * psat(T)
    return 0.622 * pv / (101325 - pv)


SURFACES = [
    # name, area, U, glazing, azimuth (deg from N, outward normal), tilt from horizontal
    ("East side panel", 1.12, 2.801, False, 90, 90),
    ("East side windows", 0.73, 2.569, True, 90, 90),
    ("East doors", 1.70, 4.89, False, 90, 90),
    ("West side panel", 1.12, 2.801, False, 270, 90),
    ("West side windows", 0.73, 2.569, True, 270, 90),
    ("West doors", 1.70, 4.89, False, 270, 90),
    ("Rear body", 0.70, 2.667, False, 180, 90),
    ("Rear window", 0.60, 2.611, True, 180, 60),
    ("Front body", 2.33, 2.667, False, 0, 90),
    ("Windshield", 1.34, 5.02, True, 0, 35),
    ("Roof", 2.00, 0.532, False, 0, 0),
    ("Floor", 6.38, 2.267, False, 0, 180),
]


def solar_on_surfaces(solar_hour=15.0):
    lat = math.radians(24.9)
    decl = math.radians(23.45)
    hra = math.radians(15 * (solar_hour - 12))
    sin_alt = math.sin(lat) * math.sin(decl) + math.cos(lat) * math.cos(decl) * math.cos(hra)
    alt = math.asin(sin_alt)
    cos_az = (math.sin(decl) - sin_alt * math.sin(lat)) / (math.cos(alt) * math.cos(lat))
    az = math.acos(max(-1, min(1, cos_az)))
    if hra > 0:
        az = 2 * math.pi - az
    # ASHRAE (1985) clear-sky constants for 21 June
    a_c, b_c, c_c = 1088.0, 0.205, 0.134
    i_dn = a_c * math.exp(-b_c / sin_alt) * A["C10"]
    i_dh = c_c * i_dn
    rho_g = 0.2
    res = {}
    for name, area, u, glz, s_az, tilt in SURFACES:
        if tilt >= 180:
            res[name] = 0.0
            continue
        tr, sa = math.radians(tilt), math.radians(s_az)
        cos_inc = math.cos(alt) * math.cos(az - sa) * math.sin(tr) + math.sin(alt) * math.cos(tr)
        direct = i_dn * max(cos_inc, 0)
        diffuse = i_dh * (1 + math.cos(tr)) / 2
        refl = rho_g * (i_dn * sin_alt + i_dh) * (1 - math.cos(tr)) / 2
        res[name] = direct + diffuse + refl
    return res, math.degrees(alt), math.degrees(az), i_dn


def cabin_heat_balance(t_out, rh_out, t_in=25.0, rh_in=50.0, solar_hour=15.0):
    flux, *_ = solar_on_surfaces(solar_hour)
    h_o = A["C11"]
    alpha = A["C12"]
    comp = {"Opaque conduction (sol-air)": 0.0, "Glazing conduction": 0.0,
            "Transmitted solar": 0.0, "Floor (road-side)": 0.0}
    for name, area, u, glz, *_ in SURFACES:
        if name == "Floor":
            comp["Floor (road-side)"] += u * area * (A["C13"] - t_in)
        elif glz:
            tau = A["C14"] if name == "Windshield" else A["C15"]
            comp["Glazing conduction"] += u * area * (t_out - t_in)
            comp["Transmitted solar"] += area * flux[name] * (tau + A["C16"] * A["C17"])
        else:
            t_sol = t_out + alpha * flux[name] / h_o
            comp["Opaque conduction (sol-air)"] += u * area * (t_sol - t_in)
    occupants = A["C18"]
    comp["Occupants sensible"] = occupants * A["C19"]
    comp["Occupants latent"] = occupants * A["C20"]
    m_vent = A["C21"] * occupants * 1.2 / 1000
    comp["Fresh-air sensible"] = m_vent * 1006 * (t_out - t_in)
    w_o, w_i = humidity_ratio(t_out, rh_out), humidity_ratio(t_in, rh_in)
    comp["Fresh-air latent"] = m_vent * 2.45e6 * max(w_o - w_i, 0)
    comp["Electronics and blower"] = A["C22"]
    return {k: v / 1000 for k, v in comp.items()}


def workbook_audit():
    # Reproduces the recovered workbook rows and the dimensionally consistent
    # recalculation used for the audit figure.
    rows = [  # name, area, U, SC, SCL, CLTDI (deg F table value), recorded Q
        ("East side panel", 1.12, 2.801, None, None, 35, 134.89),
        ("East side windows", 0.73, 2.569, 0.811, 52, None, 97.51),
        ("East doors", 1.70, 4.89, 0.811, 52, None, 432.27),
        ("West side panel", 1.12, 2.801, None, None, 20, 134.89),
        ("West side windows", 0.73, 2.569, 0.811, 112, None, 97.51),
        ("West doors", 1.70, 4.89, 0.811, 112, None, 432.27),
        ("Rear body", 0.70, 2.667, None, None, 28, 67.2),
        ("Rear window", 0.60, 2.611, 0.811, 30, None, 47.0),
        ("Front body", 2.33, 2.667, None, None, 16, 149.13),
        ("Windshield", 1.34, 5.02, 0.811, 33, None, 222.0),
        ("Roof", 2.00, 0.532, None, None, 90, 104.27),
        ("Floor", 6.38, 2.267, None, None, 90, 1417.0),
    ]
    out = []
    for name, area, u, sc, scl, cltdi, recorded in rows:
        if scl is not None:
            # Glazing: solar Q = A SC SCL (+ conduction U A dT at 38.1/23 C)
            corrected = area * sc * scl + u * area * (38.1 - 23.0)
            if "doors" in name:
                corrected = u * area * (38.1 - 23.0)  # opaque door: conduction only
        else:
            # CLTD tables are deg F differences; apply the SI correction
            cltd_k = cltdi * 5 / 9 + (25.5 - 23.0) + (38.1 - 29.4)
            corrected = u * area * cltd_k
        out.append((name, recorded, corrected))
    return out


# ---------------------------------------------------------------------------
# Figures
# ---------------------------------------------------------------------------
INK, MUTED, GRID = "#1f2328", "#5b6470", "#d9dde3"
C = ["#2563eb", "#d97706", "#059669", "#9333ea", "#dc2626", "#0891b2", "#64748b", "#ca8a04"]


def style(ax, title, xl, yl):
    ax.set_title(title, loc="left", fontsize=11, color=INK, fontweight="bold")
    ax.set_xlabel(xl, color=MUTED)
    ax.set_ylabel(yl, color=MUTED)
    ax.grid(True, color=GRID, lw=0.7)
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(colors=MUTED)


def fig_operating_points(mp, traces):
    rpm_ax, tq_ax, eta, tcurve, _ = mp
    fig, axes = plt.subplots(1, 2, figsize=(13, 5.2))
    ax = axes[0]
    rr, tt = np.meshgrid(rpm_ax, tq_ax, indexing="ij")
    cs = ax.contourf(rr, tt, eta * 100, levels=[50, 70, 80, 85, 88, 90, 92, 93, 94, 95.5], cmap="Blues", alpha=0.85)
    ax.contour(rr, tt, eta * 100, levels=[85, 90, 93], colors="white", linewidths=0.6)
    fig.colorbar(cs, ax=ax, label="Integrated efficiency (%)")
    ax.plot(tcurve[:, 0], tcurve[:, 1], color=INK, lw=1.6, label="Peak torque envelope")
    for (label, tr), col in zip(traces.items(), C):
        mot = tr["p_wheel"] > 0
        if label.startswith(("Sustained", "Low-speed")):
            ax.plot(tr["rpm"][0], tr["torque"][0], "D", ms=9, color=col, mec="white", label=label)
        else:
            ax.scatter(tr["rpm"][mot], tr["torque"][mot], s=6, color=col, alpha=0.45, label=f"{label} (1 Hz, motoring)")
    ax.set_xlim(0, 12000)
    ax.set_ylim(0, 300)
    style(ax, "Where each schedule operates on the efficiency map", "Drive-unit speed (rpm)", "Torque (Nm)")
    ax.legend(fontsize=8, loc="upper right")

    ax = axes[1]
    edges_t = np.linspace(0, 300, 13)
    edges_s = np.linspace(0, 12000, 13)
    heat = np.zeros((12, 12))
    for label in ("Urban stop-start", "Highway"):
        tr = traces[label]
        h, _, _ = np.histogram2d(tr["rpm"], np.abs(tr["torque"]), bins=[edges_s, edges_t], weights=tr["heat"] / 3600)
        heat += h
    im = ax.pcolormesh(edges_s, edges_t, heat.T * 1000, cmap="Oranges", shading="flat")
    fig.colorbar(im, ax=ax, label="Drive-unit heat (Wh)")
    ax.plot(tcurve[:, 0], tcurve[:, 1], color=INK, lw=1.6)
    ax.set_xlim(0, 12000)
    ax.set_ylim(0, 300)
    style(ax, "Heat energy by operating region (NYCC + HWFET)", "Drive-unit speed (rpm)", "|Torque| (Nm)")
    fig.suptitle("Gap fill 1: operating-point density on the supplied map (calculated, no new assumption)", x=0.01, ha="left", fontsize=12)
    fig.tight_layout()
    fig.savefig(FIG_DIR / "gap_drive_operating_points.png", dpi=150)
    plt.close(fig)


def fig_motor_calibration(mp):
    speeds, torque, e, il, ml, r_hot = implied_winding_resistance(mp)
    fig, axes = plt.subplots(1, 2, figsize=(13, 4.8))
    ax = axes[0]
    ax.plot(speeds, r_hot, color=C[0], lw=2, label="Implied winding-to-coolant R")
    ax.axhline(0.015, color=C[4], ls="--", lw=1.5, label="Configured R = 0.015 K/W")
    style(ax, "Supplier 143 C rated-rise point implies a higher resistance", "Assumed rated operating speed (rpm)", "Winding-to-coolant resistance (K/W)")
    ax.set_ylim(0, max(r_hot) * 1.2)
    ax.legend(fontsize=8)

    ax = axes[1]
    duties = {"Sustained 10% grade": 2.769, "Low-speed hot-weather grade": 1.549}
    coolant = np.linspace(45, 75, 31)
    r_mid = float(np.median(r_hot))
    for (k, q), col in zip(duties.items(), C):
        ax.plot(coolant, coolant + q * 1000 * r_mid, color=col, lw=2, label=f"{k}: implied R {r_mid:.3f} K/W (upper bound)")
        ax.plot(coolant, coolant + q * 1000 * 0.015, color=col, lw=1.2, ls="--", label=f"{k}: configured R")
    ax.axhline(180, color=INK, ls=":", lw=1.2)
    ax.text(46, 182, "Class H insulation 180 C", fontsize=8, color=INK)
    ax.axhline(150, color=MUTED, ls=":", lw=1.0)
    ax.text(46, 152, "Typical design hot-spot target 150 C", fontsize=8, color=MUTED)
    style(ax, "Steady winding hot-spot at sustained design duty", "Coolant at drive unit (C)", "Winding hot-spot (C)")
    ax.set_ylim(40, 200)
    ax.legend(fontsize=7, loc="lower right")
    fig.suptitle("Gap fill 2: winding resistance back-calculated from supplied reference (rated speed assumed)", x=0.01, ha="left", fontsize=12)
    fig.tight_layout()
    fig.savefig(FIG_DIR / "gap_motor_resistance_calibration.png", dpi=150)
    plt.close(fig)
    return speeds, ml, r_hot


def fig_radiator():
    v = np.linspace(0.5, 8, 60)
    fig, axes = plt.subplots(1, 2, figsize=(13, 4.8))
    ax = axes[0]
    for flow, col in zip([10, 20, 30], C):
        r = radiator_performance(v, flow)
        ax.plot(v, r["q"] / 1000, color=col, lw=2, label=f"{flow} L/min coolant")
    for (k, q), col in zip(DESIGN_DUTY_KW.items(), [C[4], C[3]]):
        ax.axhline(q, color=col, ls="--", lw=1.3)
        r = radiator_performance(v, 20)
        cross = np.interp(q, r["q"] / 1000, v)
        ax.plot(cross, q, "o", color=col)
        ax.annotate(f"{k}\nneeds {cross:.1f} m/s at 20 L/min", (cross, q), xytext=(max(cross - 3.2, 2.2), q + 0.15), fontsize=8, color=col)
    style(ax, "Estimated heat rejection, 65 C coolant in, 45 C air in", "Core-face air velocity (m/s)", "Heat rejection (kW)")
    ax.legend(fontsize=8, loc="lower right")

    ax = axes[1]
    r = radiator_performance(v, 20)
    ax.plot(v, r["ua"], color=C[0], lw=2, label="Estimated achieved UA (20 L/min)")
    ax.axhline(204.8, color=C[4], ls="--", lw=1.3, label="Required ideal UA, 10% grade (10 K air rise)")
    ax.axhline(111.3, color=C[3], ls="--", lw=1.3, label="Required ideal UA, low-speed grade")
    ax.axhline(665, color=MUTED, ls=":", lw=1.3, label="UA assumed in two-node model (665 W/K)")
    style(ax, "Estimated achieved UA vs requirement and model input", "Core-face air velocity (m/s)", "UA (W/K)")
    ax.legend(fontsize=8)
    fig.suptitle("Gap fill 3: candidate core performance from Chang-Wang louver correlation (louver geometry assumed)", x=0.01, ha="left", fontsize=12)
    fig.tight_layout()
    fig.savefig(FIG_DIR / "gap_radiator_performance_map.png", dpi=150)
    plt.close(fig)


def fig_battery_heat_path():
    budget = path_budget()
    r_lit = sum(budget.values())
    fig, axes = plt.subplots(1, 3, figsize=(16, 4.8))
    ax = axes[0]
    soc = np.linspace(0, 100, 201)
    i = CAPACITY_AH
    for T, col in [(25, C[0]), (45, C[1])]:
        ax.plot(soc, np.full_like(soc, i**2 * dcir(T)), color=col, lw=2, label=f"Joule, DCIR at {T} C")
    ax.plot(soc, np.full_like(soc, i**2 * ACR), color=MUTED, lw=1.2, ls="--", label="Current ACR heat floor")
    qrev = -i * (298.15) * np.interp(soc, SOC_TABLE, DUDT_TABLE)
    ax.plot(soc, qrev, color=C[2], lw=2, label="Reversible (entropic), 25 C")
    ax.axhline(0, color=INK, lw=0.6)
    style(ax, "Cell heat sources at 1C discharge", "State of charge (%)", "Cell heat (W)")
    ax.legend(fontsize=8)

    ax = axes[1]
    left = 0
    for (k, val), col in zip(budget.items(), C):
        ax.barh(["Literature build-up"], [val], left=left, color=col, label=f"{k}: {val:.3f}")
        left += val
    ax.barh(["Reconstructed path"], [R_BASE_RECON], color=C[4], alpha=0.7, label=f"Reconstructed: {R_BASE_RECON:.2f}")
    ax.set_xlim(0, 3.4)
    style(ax, f"Cell-to-coolant resistance: {r_lit:.2f} vs 3.10 K/W", "Thermal resistance (K/W)", "")
    ax.legend(fontsize=7, loc="lower right")

    ax = axes[2]
    c = np.linspace(0.1, 2, 40)
    for T_label, r_path, col, ls in [("Reconstructed 3.10 K/W, ACR", R_BASE_RECON, C[4], "--"),
                                     (f"Literature {r_lit:.2f} K/W, DCIR(25 C) + peak entropic", r_lit, C[0], "-")]:
        if "ACR" in T_label:
            q = (c * i) ** 2 * ACR
        else:
            q = (c * i) ** 2 * dcir(25) + c * i * 298.15 * 0.37e-3
        ax.plot(c, LIMIT_ABS - q * r_path, color=col, lw=2, ls=ls, label=T_label + ", 60 C limit")
        ax.plot(c, LIMIT_CHARGE - q * r_path, color=col, lw=1.2, ls=ls, alpha=0.6, label=T_label + ", 55 C limit")
    ax.axhline(45, color=MUTED, ls=":", lw=1.2)
    ax.text(0.12, 46, "45 C ambient: ambient-air radiator cannot supply below this", fontsize=7, color=MUTED)
    style(ax, "Maximum allowable coolant temperature", "Sustained C-rate", "Coolant temperature (C)")
    ax.set_ylim(-40, 65)
    ax.legend(fontsize=6.5, loc="lower left")
    fig.suptitle("Gap fill 4: battery heat terms and cell-to-coolant path from literature values", x=0.01, ha="left", fontsize=12)
    fig.tight_layout()
    fig.savefig(FIG_DIR / "gap_battery_heat_and_path.png", dpi=150)
    plt.close(fig)
    return budget, r_lit


def fig_battery_transient(r_lit):
    fig, axes = plt.subplots(1, 2, figsize=(13, 4.8))
    for ax, t_cool, title in [(axes[0], 25.0, "Chiller-conditioned coolant, 25 C"),
                              (axes[1], 50.0, "Radiator-only coolant at 45 C ambient (~50 C)")]:
        for c_rate, col in zip([0.5, 1.0, 2.0], C):
            t, T, *_ = battery_discharge(c_rate, t_cool, r_lit)
            ax.plot(t / 60, T, color=col, lw=2, label=f"{c_rate:g}C, literature path {r_lit:.2f} K/W")
            t, T, *_ = battery_discharge(c_rate, t_cool, R_BASE_RECON)
            ax.plot(t / 60, T, color=col, lw=1.2, ls="--", label=f"{c_rate:g}C, reconstructed 3.10 K/W")
        ax.axhline(LIMIT_ABS, color=C[4], ls=":", lw=1.3)
        ax.axhline(LIMIT_CHARGE, color=C[1], ls=":", lw=1.3)
        ax.text(1, LIMIT_ABS + 0.6, "60 C absolute", fontsize=8, color=C[4])
        ax.text(1, LIMIT_CHARGE + 0.6, "55 C charge cutoff", fontsize=8, color=C[1])
        style(ax, title, "Time from full charge (min)", "Cell temperature (C)")
        ax.set_ylim(t_cool - 2, max(75, t_cool + 30))
        ax.legend(fontsize=7, loc="upper right")
    fig.suptitle("Gap fill 5: constant-current discharge transient (single lumped cell, DCIR + entropic heat)", x=0.01, ha="left", fontsize=12)
    fig.tight_layout()
    fig.savefig(FIG_DIR / "gap_battery_discharge_transient.png", dpi=150)
    plt.close(fig)


def fig_cabin():
    audit = workbook_audit()
    scen = {"Workbook basis\n38.1 C, as recovered": None,
            "Dry heat\n45 C, 25% RH": cabin_heat_balance(45, 25),
            "Humid heat (2015 peak)\n45 C, 44% RH": cabin_heat_balance(45, 44)}
    fig, axes = plt.subplots(1, 3, figsize=(17, 5.2), gridspec_kw=dict(width_ratios=[1.1, 1.3, 1]))
    ax = axes[0]
    names = [a[0] for a in audit]
    y = np.arange(len(names))
    ax.barh(y + 0.2, [a[1] for a in audit], height=0.4, color=C[4], label=f"Recorded in workbook ({sum(a[1] for a in audit):.0f} W)")
    ax.barh(y - 0.2, [a[2] for a in audit], height=0.4, color=C[0], label=f"Recomputed, same inputs ({sum(a[2] for a in audit):.0f} W)")
    ax.set_yticks(y, names, fontsize=8)
    ax.invert_yaxis()
    style(ax, "Workbook audit: body and glazing rows", "Load (W)", "")
    ax.legend(fontsize=8, loc="center right")

    ax = axes[1]
    labels = list(scen.keys())
    comps = list(scen[labels[1]].keys())
    recovered = {"Body and glazing (as recorded)": 3.336, "Occupants (as recorded)": 0.594, "Infiltration (as recorded)": 0.226}
    bottom = 0
    for (k, v), col in zip(recovered.items(), [C[6], "#94a3b8", "#cbd5e1"]):
        ax.bar(0, v, bottom=bottom, color=col, label=k)
        bottom += v
    ax.text(0, bottom + 0.1, f"{bottom:.2f} kW", ha="center", fontsize=9)
    palette = ["#2563eb", "#60a5fa", "#d97706", "#a16207", "#059669", "#6ee7b7", "#9333ea", "#c084fc", "#64748b"]
    for xi, lab in enumerate(labels[1:], start=1):
        bottom = 0
        for comp, col in zip(comps, palette):
            ax.bar(xi, scen[lab][comp], bottom=bottom, color=col, label=comp if xi == 1 else None)
            bottom += scen[lab][comp]
        ax.text(xi, bottom + 0.1, f"{bottom:.2f} kW", ha="center", fontsize=9)
    ax.set_xticks(range(len(labels)), labels, fontsize=8)
    style(ax, "Steady cabin load at 15:00, 25 C / 50% RH cabin", "", "Load (kW)")
    ax.legend(fontsize=7, loc="upper left", ncol=2)
    ax.set_ylim(0, 8.5)

    ax = axes[2]
    steady = sum(scen[labels[2]].values())
    minutes = np.linspace(10, 60, 51)
    for cap, col in zip([A["C23"], A["C24"], A["C25"]], C):
        extra = cap * 1000 * (A["C26"] - 25.0) / (minutes * 60) / 1000
        ax.plot(minutes, steady + extra, color=col, lw=2, label=f"Effective interior mass {cap:.0f} kJ/K")
    ax.axvline(30, color=MUTED, ls=":", lw=1.2)
    ax.text(30.5, ax.get_ylim()[1] * 0.95 if ax.get_ylim()[1] > 0 else 10, "30 min target", fontsize=8, color=MUTED)
    style(ax, "Average capacity to pull down 80 C soak to 25 C", "Pull-down time (min)", "Mean cooling capacity (kW)")
    ax.legend(fontsize=8)
    fig.suptitle("Gap fill 6: cabin workbook audit and heat-balance rebuild (Fayazbakhsh and Bahrami structure)", x=0.01, ha="left", fontsize=12)
    fig.tight_layout()
    fig.savefig(FIG_DIR / "gap_cabin_heat_balance.png", dpi=150)
    plt.close(fig)
    return audit, scen


def main():
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    mp = load_map()
    traces = drive_traces(mp)
    for k, tr in traces.items():
        print(f"{k}: mean heat {np.trapezoid(tr['heat'], tr['t']) / (tr['t'][-1] - tr['t'][0]):.3f} kW")
    fig_operating_points(mp, traces)
    speeds, ml, r_hot = fig_motor_calibration(mp)
    print("Implied R range", r_hot.min(), r_hot.max(), "median", np.median(r_hot))
    print("motor-only loss range kW", ml.min(), ml.max())
    fig_radiator()
    for v in (2.97, 4.0, 6.0):
        r = radiator_performance(np.array([v]), 20)
        print(f"Radiator at {v} m/s: Q {r['q'][0]/1000:.3f} kW UA {r['ua'][0]:.1f} h_air {r['h_air'][0]:.1f} "
              f"eta_f {r['eta_fin'][0]:.3f} ReLp {r['re_lp'][0]:.0f} Re_c {r['re_c']:.0f} h_c {r['h_c']:.0f} dp {r['dp']/1000:.3f} kPa sigma {r['sigma']:.3f}")
    budget, r_lit = fig_battery_heat_path()
    print("Battery path", budget, r_lit)
    print("DCIR25", dcir(25), "DCIR45", dcir(45), "DCIR0", dcir(0))
    fig_battery_transient(r_lit)
    for cr in (0.5, 1, 2):
        for tc in (25, 50):
            _, T, *_ = battery_discharge(cr, tc, r_lit)
            _, T2, *_ = battery_discharge(cr, tc, R_BASE_RECON)
            print(f"{cr}C coolant {tc}: peak lit {T.max():.2f}  recon {T2.max():.2f}")
    audit, scen = fig_cabin()
    print("Audit recorded", sum(a[1] for a in audit), "recomputed", sum(a[2] for a in audit))
    for k, v in scen.items():
        if v:
            print(k.replace("\n", " "), {kk: round(vv, 3) for kk, vv in v.items()}, "total", round(sum(v.values()), 3))
    print(solar_on_surfaces())


if __name__ == "__main__":
    main()
